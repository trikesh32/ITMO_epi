#!/usr/bin/env python3
"""Проверка поставляемых артефактов и независимого пересчёта Excel."""
from pathlib import Path
import json, math, re, sys
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from data import TASKS, REQS, UCS, AUTHORS
root=Path(__file__).resolve().parents[1]
r=json.loads((root/'docs/calculations.json').read_text())
book=next(root.glob('*.xlsx'))
w=load_workbook(book,data_only=False)
c=load_workbook(book,data_only=True)
errors=[(s.title,x.coordinate,x.value) for s in c for row in s for x in row if x.data_type=='e']
assert not errors,errors
mapping={2:'naive_hours',3:'pert_hours',4:'pert_sigma',7:'critical_hours',8:'ilf',10:'transactions_fp',11:'ufp',13:'afp',15:'ksloc',19:'cocomo_pm',20:'cocomo_hours',22:'cocomo_months',24:'uucw',25:'uaw',27:'tcf',29:'ecf',30:'ucp',32:'ucp_hours_pf20'}
cached=c['Итоги']['B2'].value is not None
if cached:
 for row,key in mapping.items():
  actual=c['Итоги'].cell(row,2).value
  assert actual is not None and math.isclose(actual,r[key],rel_tol=1e-9),(row,key,actual,r[key])
else:
 print('Кэш формул отсутствует: откройте и сохраните книгу в Excel/LibreOffice для независимой проверки значений.')
 if '--require-cache' in sys.argv: raise SystemExit(1)
assert len(w['Работы']['A'])-2==len(TASKS)==32
assert len(w['Требования']['A'])-1==len(REQS)==24
assert len(w['UCP прецеденты']['A'])-1==len(UCS)==25
network=ET.parse(root/'Сетевая диаграмма.drawio').getroot()
actual={(e.get('source'),e.get('target')) for e in network.iter('mxCell') if e.get('edge')=='1'}
expected={(d,t['id']) for t in TASKS for d in t['deps']}
assert expected==actual
for i,t in enumerate(TASKS,2):
 assert w['CPM'].cell(i,3).data_type=='f'
 if cached:
  assert math.isclose(c['CPM'].cell(i,3).value,(t['o']+4*t['m']+t['p'])/6)
# Ресурсные конфликты и выполнение всех зависимостей.
plan={row[0]:(row[2],row[3],row[4]) for row in c['Расписание'].iter_rows(min_row=2,values_only=True)}
for t in TASKS:
 worker,start,end=plan[t['id']]
 assert math.isclose(end-start,(t['o']+4*t['m']+t['p'])/6,abs_tol=1e-8)
 assert all(plan[d][2]<=start+1e-8 for d in t['deps'])
for worker in range(1,r['team']+1):
 intervals=sorted((s,e) for w,s,e in plan.values() if w==worker)
 assert all(a[1]<=b[0]+1e-8 for a,b in zip(intervals,intervals[1:]))
assert math.isclose(max(p[2] for p in plan.values()),r['schedule_hours'])
pdf=PdfReader(root/'docs/report.pdf'); text='\n'.join(p.extract_text() or '' for p in pdf.pages)
for name in AUTHORS: assert name in text,name
for req in REQS: assert req[0] in text,req[0]
assert 'P3406' in text and 'stepik.org' in text
assert 'калибровка PF' in text and r['pf_calibrated'] is False
assert 'KudaGo' not in text and 'Барсуков' not in text
assert len(pdf.pages)>10
for path in [root/'README.md',root/'docs/report.md']:
 for link in re.findall(r'\]\(([^)]+)\)',path.read_text()):
  if '://' not in link and not link.startswith('#'):
   from urllib.parse import unquote
   assert (path.parent/unquote(link)).exists(),(path,link)
print(f'OK: {len(pdf.pages)} страниц PDF, 24 FR, 32 работы, {len(expected)} зависимостей, 25 UC.')
print(f'Excel: {len(w.sheetnames)} листов; ошибок формул нет; кэш проверен: {cached}.')
