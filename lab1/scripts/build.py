#!/usr/bin/env python3
"""Генерация отчёта, формульной книги и диаграмм из единого набора данных."""
from pathlib import Path
import json, math, textwrap, html, xml.etree.ElementTree as ET
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
import pymupdf as fitz
from data import *
ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT/'docs'
DOCS.mkdir(exist_ok=True)
FONT = '/System/Library/Fonts/Supplemental/Arial.ttf'
if not Path(FONT).exists():
 FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
def fmt(x):
 if isinstance(x,float): return f'{x:,.2f}'.replace(',',' ').replace('.',',')
 return str(x)
def fcomplex(kind,ret,det):
 if kind in ('ILF','EIF'):
  r=0 if ret==1 else 1 if ret<=5 else 2
  c=0 if det<=19 else 1 if det<=50 else 2
  return [[0,0,1],[0,1,2],[1,2,2]][r][c]
 if kind=='EI':
  r=0 if ret<=1 else 1 if ret==2 else 2
  c=0 if det<=4 else 1 if det<=15 else 2
 else:
  r=0 if ret<=1 else 1 if ret<=3 else 2
  c=0 if det<=5 else 1 if det<=19 else 2
 return [[0,0,1],[0,1,2],[1,2,2]][r][c]
weights={'ILF':[7,10,15],'EIF':[5,7,10],'EI':[3,4,6],'EO':[4,5,7],'EQ':[3,4,6]}
labels=['Низкая','Средняя','Высокая']
byid={t['id']:t for t in TASKS}
succ=defaultdict(list)
for t in TASKS:
 assert t['o']<=t['m']<=t['p']
 t['te']=(t['o']+4*t['m']+t['p'])/6
 t['var']=((t['p']-t['o'])/6)**2
 t['es']=max((byid[d]['ef'] for d in t['deps']),default=0)
 t['ef']=t['es']+t['te']
 for d in t['deps']: succ[d].append(t['id'])
CP=max(t['ef'] for t in TASKS)
for t in reversed(TASKS):
 t['lf']=min((byid[s]['ls'] for s in succ[t['id']]),default=CP)
 t['ls']=t['lf']-t['te']; t['slack']=t['ls']-t['es']
 t['bottom']=t['te']+max((byid[s]['bottom'] for s in succ[t['id']]),default=0)
critical=[t['id'] for t in TASKS if abs(t['slack'])<1e-8]
# В данном графе критический путь единственный; проверяем последовательность.
assert all(a in byid[b]['deps'] for a,b in zip(critical,critical[1:]))
EFF=sum(t['te'] for t in TASKS); NAIVE=sum(t['m'] for t in TASKS)
SIG=math.sqrt(sum(t['var'] for t in TASKS)); CPSIG=math.sqrt(sum(byid[i]['var'] for i in critical))
def schedule(n):
 remaining=set(byid); done=set(); running=[]; time=0.; result=[]; free=list(range(1,n+1))
 while remaining or running:
  ready=sorted((byid[i] for i in remaining if set(byid[i]['deps'])<=done),key=lambda t:(-t['bottom'],t['id']))
  while free and ready:
   t=ready.pop(0); worker=free.pop(0); end=time+t['te']
   running.append((end,t['id'],worker)); remaining.remove(t['id']); result.append((t['id'],worker,time,end))
  if running:
   time=min(x[0] for x in running)
   completed=[x for x in running if abs(x[0]-time)<1e-8]
   for x in completed: running.remove(x); done.add(x[1]); free.append(x[2])
   free.sort()
  else: assert not remaining
 return time,result
SCHEDULES={n:schedule(n) for n in range(1,9)}
TEAM=min(n for n,(d,_) in SCHEDULES.items() if d<=1.10*CP)
DURATION,PLAN=SCHEDULES[TEAM]
# Верификация выполнимости расписания.
planmap={i:(w,s,e) for i,w,s,e in PLAN}
for i,w,s,e in PLAN:
 assert all(planmap[d][2]<=s+1e-8 for d in byid[i]['deps'])
for worker in range(1,TEAM+1):
 events=sorted((s,e) for i,w,s,e in PLAN if w==worker)
 assert all(a[1]<=b[0]+1e-8 for a,b in zip(events,events[1:]))
FPROWS=[]
for i,n,k,r,d,x in FILES:
 c=fcomplex(k,r,d); FPROWS.append([i,n,k,r,d,labels[c],weights[k][c],x])
TXROWS=[]
for t in TX:
 c=fcomplex(t['type'],len(t['files']),t['det']); t['fp']=weights[t['type']][c]
 TXROWS.append([t['id'],t['name'],t['type'],len(t['files']),t['det'],labels[c],t['fp'],', '.join(t['files']),t['req']])
ILF=sum(r[6] for r in FPROWS); TRANS=sum(t['fp'] for t in TX); UFP=ILF+TRANS
VAF=.65+.01*sum(r[1] for r in GSC); AFP=UFP*VAF
SF_SUM=sum(r[1] for r in SF); E=.91+.01*SF_SUM; EAF=math.prod(r[1] for r in EM)
KSLOC=UFP*47/1000; PM=2.94*KSLOC**E*EAF; F=.28+.2*(E-.91); TDEV=3.67*PM**F
COCOMO_H=PM*152
for u in UCS: u['weight']=5 if len(u['steps'])<=3 else 10 if len(u['steps'])<=7 else 15
UUCW=sum(u['weight'] for u in UCS); UAW=sum(a[1] for a in ACTORS)
TF=sum(r[2]*r[3] for r in TCF); EF=sum(r[2]*r[3] for r in ECF)
TC=.6+.01*TF; EC=1.4-.03*EF; UCP=(UUCW+UAW)*TC*EC
assert {r[0] for r in REQS}==set(' '.join(t['req'] for t in TX).split())
assert {r[0] for r in REQS}==set(' '.join(u['req'] for u in UCS).split())

# Диаграмма AON: все зависимости, единые координаты для draw.io, SVG и PNG.
levels={}
for t in TASKS: levels[t['id']]=max((levels[d]+1 for d in t['deps']),default=0)
layer=defaultdict(list)
for i,l in levels.items(): layer[l].append(i)
W=1500; H=(max(layer)+1)*155+115
positions={}
for level,ids in layer.items():
 for j,i in enumerate(ids): positions[i]=(W/2+(j-(len(ids)-1)/2)*285-130,80+level*155)
def draw_diagram():
 root=ET.Element('mxfile',host='app.diagrams.net',version='24.7.17')
 dg=ET.SubElement(root,'diagram',id='stepik-network',name='PERT — Stepik')
 model=ET.SubElement(dg,'mxGraphModel',dx=str(W),dy=str(H),grid='1',page='1',pageWidth=str(W),pageHeight=str(H))
 cells=ET.SubElement(model,'root'); ET.SubElement(cells,'mxCell',id='0'); ET.SubElement(cells,'mxCell',id='1',parent='0')
 sv=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">', '<defs><marker id="arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto"><path d="M0,0 L8,4 L0,8" fill="#64748b"/></marker></defs>', '<rect width="100%" height="100%" fill="white"/>']
 pdf=fitz.open(); page=pdf.new_page(width=W,height=H); page.insert_font(fontname='ru',fontfile=FONT)
 title='Stepik: сетевая диаграмма PERT. Красный — критический путь. Время: рабочие часы.'
 page.insert_text((35,35),title,fontname='ru',fontsize=21)
 sv.append(f'<text x="35" y="35" font-family="Arial" font-size="21">{title}</text>')
 for t in TASKS:
  x,y=positions[t['id']]
  for d in t['deps']:
   dx,dy=positions[d]; crit=d in critical and t['id'] in critical and abs(byid[d]['ef']-t['es'])<1e-8
   color='#c0392b' if crit else '#aab4c0'; rgb=(.75,.22,.17) if crit else (.67,.70,.75)
   pts=[(dx+130,dy+96),(dx+130,dy+120),(x+130,y-20),(x+130,y)]
   for a,b in zip(pts,pts[1:]): page.draw_line(a,b,color=rgb,width=2 if crit else 1)
   page.draw_line((x+125,y-8),(x+130,y),color=rgb); page.draw_line((x+135,y-8),(x+130,y),color=rgb)
   sv.append(f'<polyline points="'+ ' '.join(f'{a},{b}' for a,b in pts)+f'" fill="none" stroke="{color}" stroke-width="2" marker-end="url(#arrow)"/>')
   e=ET.SubElement(cells,'mxCell',id=f'{d}-{t["id"]}',edge='1',parent='1',source=d,target=t['id'],style=f'edgeStyle=orthogonalEdgeStyle;endArrow=block;strokeColor={color};strokeWidth=2;')
   ET.SubElement(e,'mxGeometry',relative='1',attrib={'as':'geometry'})
 for t in TASKS:
  i=t['id']; x,y=positions[i]; crit=i in critical; fill='#fde9e7' if crit else '#edf4fb'
  label=[f'{i} · {fmt(t["te"])} ч']+textwrap.wrap(t['name'],29)+[f'ES {fmt(t["es"])} · резерв {fmt(max(0,t["slack"]))}']
  page.draw_rect(fitz.Rect(x,y,x+260,y+96),color=(.7,.2,.15) if crit else (.3,.4,.55),fill=(.99,.92,.9) if crit else (.93,.96,.99))
  sv.append(f'<rect x="{x}" y="{y}" width="260" height="96" rx="6" fill="{fill}" stroke="#64748b"/>')
  for k,line in enumerate(label):
   page.insert_text((x+10,y+20+k*18),line,fontname='ru',fontsize=13)
   sv.append(f'<text x="{x+10}" y="{y+20+k*18}" font-family="Arial" font-size="13">{html.escape(line)}</text>')
  c=ET.SubElement(cells,'mxCell',id=i,value='\n'.join(label),vertex='1',parent='1',style=f'rounded=1;whiteSpace=wrap;html=0;fillColor={fill};strokeColor=#64748b;fontSize=13;')
  ET.SubElement(c,'mxGeometry',x=str(x),y=str(y),width='260',height='96',attrib={'as':'geometry'})
 ET.indent(root); ET.ElementTree(root).write(ROOT/'Сетевая диаграмма.drawio',encoding='utf-8',xml_declaration=True)
 (DOCS/'network.svg').write_text('\n'.join(sv+['</svg>']))
 page.get_pixmap(matrix=fitz.Matrix(1,1)).save(DOCS/'network.png')
 pdf.save(DOCS/'network.pdf')
 # Гант по каждому исполнителю, с подписями ID.
 gp=fitz.open(); pg=gp.new_page(width=1400,height=190+TEAM*95); pg.insert_font(fontname='ru',fontfile=FONT)
 pg.insert_text((35,35),f'Ресурсный план: {TEAM} исполнителя; {fmt(DURATION/40)} рабочих недель',fontname='ru',fontsize=23)
 scale=1200/DURATION
 for tick in range(0,math.ceil(DURATION/40)+1,5):
  x=140+tick*40*scale; pg.draw_line((x,70),(x,110+TEAM*95),color=(.85,.85,.85)); pg.insert_text((x,65),str(tick),fontsize=12)
 for worker in range(1,TEAM+1):
  y=90+(worker-1)*95; pg.insert_text((20,y+30),f'Участник {worker}',fontname='ru',fontsize=14)
  for i,w,s,e in PLAN:
   if w!=worker: continue
   rect=fitz.Rect(140+s*scale,y,140+e*scale,y+48)
   pg.draw_rect(rect,color=(.25,.4,.55),fill=(.84,.92,.98))
   pg.insert_text((rect.x0+2,y+25),i,fontname='ru',fontsize=9)
 pg.insert_text((35,145+TEAM*95),'Ось X: рабочие недели по 40 часов. Внутри полос — идентификаторы работ.',fontname='ru',fontsize=14)
 pg.get_pixmap().save(DOCS/'schedule.png')
 gp.close(); pdf.close()
draw_diagram()

# Табличная книга: формулы без ручного дублирования итогов.
wb=Workbook(); wb.remove(wb.active); wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
def sheet(name,headers,rows,widths=None):
 ws=wb.create_sheet(name); ws.append(headers)
 for row in rows: ws.append(row)
 ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
 ws.sheet_view.showGridLines=False
 for c in ws[1]: c.font=Font(name='Calibri',bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='24476B')
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True); c.font=Font(name='Calibri',size=11)
   if c.row%2==0: c.fill=PatternFill('solid',fgColor='EFF4F9')
   if c.data_type=='f' or isinstance(c.value,float): c.number_format='0.00'
 for j in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(j)].width=(widths[j-1] if widths else 22)
 ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:1'
 return ws
sheet('Пояснения',['Параметр','Значение'],[
 ['Вариант','Stepik.org; учебная оценка разработки аналога в оговорённых границах'],['Авторы','; '.join(AUTHORS)+'; '+GROUP],['Дата',DATE],
 ['Единицы','O/M/P/TE — чел·ч. CPM — рабочие часы при одном исполнителе на работу. Неделя = 40 часов.'],
 ['Границы','Веб-версия, 24 FR; внешние готовые платёжный, grader, email и media API.'],
 ['Происхождение оценок','Экспертные учебные допущения, не трудозатраты реальной команды Stepik.'],
 ['Методы','Наивный = сумма M; PERT = сумма TE; критический путь не является трудоёмкостью.'],
 ['Ресурсный план','Эвристика приоритета по длине оставшегося пути; 1 исполнитель на задачу, без прерывания.'],
 ['Рекомендуемая команда',TEAM],['Критерий','Наименьшая из построенных команд 1–8 со сроком не более 110% CPM. Глобальный оптимум не доказывается.'],
 ['Функциональные точки','DET/RET/FTR — предварительная оценка. EIF=0: внешние API не считаются файлами автоматически.'],
 ['COCOMO II','Early Design 2000; в SLOC переводится UFP, не AFP. JavaScript: 47 SLOC/FP.'],
 ['UCP / PF','Завершённого проекта с известными часами у авторов нет. PF=10/20/30 — сценарии, не калибровка.'],
 ['Пересчёт','При изменении входных данных в data.py запустить build.py; Excel пересчитывает формулы, диаграммы и ресурсный план пересоздаются скриптом.'],
 ],[27,120])
sheet('Требования',['ID','Блок','Функция','Критерий приёмки','Источник / допущение'],REQS,[12,22,36,95,38])
rows=[]
index={t['id']:j+2 for j,t in enumerate(TASKS)}
for t in TASKS:
 r=index[t['id']]
 rows.append([t['id'],t['name'],t['detail'],t['req'],', '.join(t['deps']),t['o'],t['m'],t['p'],f'=(F{r}+4*G{r}+H{r})/6',f'=(H{r}-F{r})/6',f'=J{r}^2'])
ws=sheet('Работы',['ID','Работа','Результат','FR','Предшественники','O, ч','M, ч','P, ч','TE, ч','σ, ч','Дисперсия, ч²'],rows,[10,38,75,28,30,12,12,12,14,14,18])
end=len(TASKS)+1; total=end+1
ws.append(['ИТОГО','','','','',f'=SUM(F2:F{end})',f'=SUM(G2:G{end})',f'=SUM(H2:H{end})',f'=SUM(I2:I{end})',f'=SQRT(SUM(K2:K{end}))',f'=SUM(K2:K{end})'])
rows=[]
for t in TASKS:
 r=index[t['id']]; dep=','.join(f'E{index[d]}' for d in t['deps'])
 following=','.join(f'F{index[s]}' for s in succ[t['id']])
 rows.append([t['id'],t['name'],f"='Работы'.I{r}".replace("'.","'!"),f'=MAX({dep})' if dep else '=0',f'=D{r}+C{r}',f'=G{r}-C{r}',f'=MIN({following})' if following else f'=MAX(E2:E{end})',f'=F{r}-D{r}',f'=IF(ABS(H{r})<0.000001,1,0)'])
sheet('CPM',['ID','Работа','TE','ES','EF','LS','LF','Резерв','Критическая'],rows,[10,42,14,14,14,14,14,14,14])
sheet('Команда',['Исполнителей','Нижняя граница, ч','Расписание, ч','Рабочих недель','Загрузка','Отклонение от CPM'],[[n,max(CP,EFF/n),d,d/40,EFF/(n*d),d/CP-1] for n,(d,_) in SCHEDULES.items()],[18,24,24,22,18,26])
sheet('Расписание',['ID','Работа','Исполнитель','Старт, ч','Финиш, ч','Старт, неделя','Финиш, неделя'],[[i,byid[i]['name'],w,s,e,s/40,e/40] for i,w,s,e in sorted(PLAN)],[12,44,18,18,18,22,22])
fws=sheet('FP данные',['ID','Логическая группа','Тип','RET','DET','Сложность','FP','Состав'],FPROWS,[12,34,12,12,12,18,12,85])
tws=sheet('FP транзакции',['ID','Процесс','Тип','FTR','DET','Сложность','FP','Группы данных','FR'],TXROWS,[12,48,12,12,12,18,12,35,22])
# Формулы сложности учитывают матрицы порогов.
for ws,kind in [(fws,'data'),(tws,'transaction')]:
 for r in range(2,ws.max_row+1):
  if kind=='data':
   expr=f'IF(D{r}=1,IF(E{r}<=50,1,2),IF(D{r}<=5,IF(E{r}<=19,1,IF(E{r}<=50,2,3)),IF(E{r}<=19,2,3)))'
   ws.cell(r,7,f'=CHOOSE({expr},7,10,15)')
  else:
   ei=f'IF(D{r}<=1,IF(E{r}<=15,1,2),IF(D{r}=2,IF(E{r}<=4,1,IF(E{r}<=15,2,3)),IF(E{r}<=4,2,3)))'
   eq=f'IF(D{r}<=1,IF(E{r}<=19,1,2),IF(D{r}<=3,IF(E{r}<=5,1,IF(E{r}<=19,2,3)),IF(E{r}<=5,2,3)))'
   expr=f'IF(C{r}="EI",{ei},{eq})'
   ws.cell(r,7,f'=IF(C{r}="EI",CHOOSE({expr},3,4,6),IF(C{r}="EO",CHOOSE({expr},4,5,7),CHOOSE({expr},3,4,6)))')
  ws.cell(r,6,f'=CHOOSE({expr},"Низкая","Средняя","Высокая")')
sheet('VAF',['Характеристика','DI (0–5)'],GSC,[48,18])
sheet('COCOMO факторы',['Фактор','Значение','Уровень','Обоснование'],SF+EM,[15,18,20,100])
sheet('UCP прецеденты',['ID','Прецедент','Акторы','FR','Транзакции сценария','Количество','Вес'],[[u['id'],u['name'],u['actor'],u['req'],'; '.join(f'{j+1}. {s}' for j,s in enumerate(u['steps'])),len(u['steps']),f'=IF(F{r}<=3,5,IF(F{r}<=7,10,15))'] for r,u in enumerate(UCS,2)],[12,40,42,20,115,18,12])
sheet('UCP акторы',['Актор','Вес','Интерфейс'],ACTORS,[28,14,70])
sheet('UCP факторы',['ID','Фактор','Вес','Оценка','Произведение','Обоснование'],[[i,n,w,v,f'=C{r}*D{r}',x] for r,(i,n,w,v,x) in enumerate(TCF+ECF,2)],[12,40,12,14,20,70])
summary=[
 ['Наивная трудоёмкость, чел·ч',f"='Работы'!G{total}"], #2
 ['PERT трудоёмкость, чел·ч',f"='Работы'!I{total}"], #3
 ['PERT σ, чел·ч',f"='Работы'!J{total}"], #4
 ['95% нижняя (приближение)','=B3-1.96*B4'],['95% верхняя (приближение)','=B3+1.96*B4'],
 ['CPM, рабочие часы',f"=MAX('CPM'!E2:E{end})"], #7
 ['ILF',f"=SUM('FP данные'!G2:G{len(FILES)+1})"],['EIF',0],['Транзакции',f"=SUM('FP транзакции'!G2:G{len(TX)+1})"],
 ['UFP','=SUM(B8:B10)'],['VAF',"=0.65+0.01*SUM('VAF'!B2:B15)"],['AFP','=B11*B12'], #13
 ['SLOC / UFP',47],['KSLOC','=B11*B14/1000'],['ΣSF',"=SUM('COCOMO факторы'!B2:B6)"],['E','=0.91+0.01*B16'],['EAF',"=PRODUCT('COCOMO факторы'!B7:B13)"],
 ['COCOMO PM','=2.94*B15^B17*B18'],['COCOMO чел·ч','=B19*152'],['F','=0.28+0.2*(B17-0.91)'],['TDEV, месяцев','=3.67*B19^B21'],['Средняя команда COCOMO','=B19/B22'],
 ['UUCW',f"=SUM('UCP прецеденты'!G2:G{len(UCS)+1})"],['UAW',"=SUM('UCP акторы'!B2:B9)"],['TF',"=SUM('UCP факторы'!E2:E14)"],['TCF','=0.6+0.01*B26'],['EF',"=SUM('UCP факторы'!E15:E22)"],['ECF','=1.4-0.03*B28'],['UCP','=(B24+B25)*B27*B29'],
 ['PF — сценарий, не калибровка',20],['UCP трудоёмкость, сценарий','=B30*B31'],
 ]
sheet('Итоги',['Показатель','Значение'],summary,[48,32])
sheet('PF чувствительность',['PF, чел·ч/UCP','Трудоёмкость, чел·ч','Статус'],[[p,f"='Итоги'!B30*A{r}",'Сценарий; данных завершённого проекта нет'] for r,p in enumerate([10,20,30],2)],[25,28,70])
sheet('Источники',['ID','Название','URL','Что подтверждает'],SOURCES,[12,60,100,95])
wb.save(ROOT/'Набор функциональных требований для разработки проекта.xlsx')

# Отчёт DOCX + текстовая версия: один источник содержания.
doc=Document(); sec=doc.sections[0]; sec.page_height=Cm(29.7); sec.page_width=Cm(21)
sec.top_margin=Cm(1.7); sec.bottom_margin=Cm(1.7); sec.left_margin=Cm(2); sec.right_margin=Cm(1.5)
normal=doc.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
for name in ['Heading 1','Heading 2','Heading 3']:
 doc.styles[name].font.name='Times New Roman'; doc.styles[name].font.color.rgb=RGBColor.from_string('203E5A')
 doc.styles[name].paragraph_format.keep_with_next=True
for name,size in [('Heading 1',16),('Heading 2',13),('Heading 3',11)]: doc.styles[name].font.size=Pt(size)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
md=[]
def p(s): doc.add_paragraph(s); md.append(s+'\n')
def h(s,level=1,page=False):
 if page: doc.add_page_break()
 doc.add_heading(s,level); md.append('#'*(level+1)+' '+s+'\n')
def table(headers,rows,widths=None):
 if widths is None:
  scores=[]
  for j,title in enumerate(headers):
   vals=[str(row[j]) for row in rows]
   if title=='ID': score=1.0
   elif all(isinstance(row[j],(int,float)) for row in rows): score=1.5
   else: score=max(1.6,min(7.5,max([len(title)]+[len(v) for v in vals])*0.095))
   scores.append(score)
  widths=[17.5*x/sum(scores) for x in scores]
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False
 for col,w in zip(t.columns,widths): col.width=Cm(w)
 for c,v in zip(t.rows[0].cells,headers): c.text=str(v)
 repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for vals in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,vals): c.text=fmt(v)
 for idx,row in enumerate(t.rows):
  cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
  for j,c in enumerate(row.cells):
   c.width=Cm(widths[j])
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3); par.paragraph_format.space_before=Pt(2); par.paragraph_format.line_spacing=1.0
    if len(rows)<=8 and idx<len(rows): par.paragraph_format.keep_with_next=True
    for run in par.runs: run.font.size=Pt(9 if len(headers)>5 else 10)
   if idx==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E7EEF5'); c._tc.get_or_add_tcPr().append(sh)
    for run in c.paragraphs[0].runs: run.bold=True
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |')
 for row in rows: md.append('| '+' | '.join(fmt(v).replace('|','/').replace('\n','; ') for v in row)+' |')
 md.append('')
def picture(filename,caption,width=17):
 doc.add_picture(str(DOCS/filename),width=Cm(width)); doc.add_paragraph(caption,style='Caption'); md.append(f'![{caption}]({filename})\n')
# Титульный лист.
for s in ['Министерство науки и высшего образования Российской Федерации','Федеральное государственное автономное образовательное учреждение высшего образования','«Национальный исследовательский университет ИТМО»','Факультет программной инженерии и компьютерной техники']:
 par=doc.add_paragraph(s); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3): doc.add_paragraph('')
par=doc.add_paragraph('ЛАБОРАТОРНАЯ РАБОТА № 1'); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.runs[0].bold=True; par.runs[0].font.size=Pt(18)
for s in ['по дисциплине «Экономика программной инженерии»','Оценка трудоёмкости разработки веб-проекта','Вариант: stepik.org']:
 par=doc.add_paragraph(s); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
for _ in range(3): doc.add_paragraph('')
for s in ['Работу выполнили:',*AUTHORS,'Группа '+GROUP]:
 par=doc.add_paragraph(s); par.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for _ in range(3): doc.add_paragraph('')
par=doc.add_paragraph('Санкт-Петербург\n2026'); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
md += ['# Лабораторная работа № 1 — Stepik\n','Авторы: '+', '.join(AUTHORS)+'. Группа '+GROUP+'.\n']
h('Содержание',page=True)
for s in ['1. Задание и границы исследования','2. Функциональные требования','3. Наивная оценка','4. PERT, критический путь и команда','5. Функциональные точки','6. COCOMO II','7. Use Case Points','8. Сравнение результатов и вывод','9. Источники']:
 doc.add_paragraph(s); md.append(s)
h('1. Задание и границы исследования',page=True)
p('Цель работы — сформировать требования к веб-проекту по варианту stepik.org и оценить разработку несколькими методами. Объект оценки — новый веб-продукт с согласованным набором функций аналога Stepik. Числа не описывают фактическую стоимость, штат или историю разработки Stepik.')
p('Задание: сформировать функциональные требования; выполнить наивную оценку и PERT; построить сетевую диаграмму, найти критический путь и предложить команду; оценить функциональные точки и трудоёмкость COCOMO II; выполнить UCP с калибровкой PF по собственному завершённому проекту; сравнить результаты. Planning Poker требуется для трёх участников; работу выполняют два человека.')
p('Исследование проведено '+DATE+' по публичному каталогу [S1] и официальным справочным статьям [S2–S11]. Каталог загружается динамически; его интерфейс и закрытые кабинеты не проходили интерактивное тестирование. Требования основаны на документации. Конкретные правила приёмки, архитектура, нагрузка и численные оценки являются проектными допущениями.')
p('Веб-версия включает аккаунты, каталог, обучение, тесты трёх типов, задачи на код, ручную проверку преподавателем, прогресс, сертификаты, обсуждения, авторский редактор, отчёты, оплату, промокоды и модерацию. Структура курса и создание материалов подтверждены [S3, S4], проверка заданий — [S5], покупка доступа — [S6].')
p('Внешние готовые сервисы: один платёжный провайдер, изолированная проверка Python-кода, доставка email и объектное хранилище видео/файлов. Включены разработка адаптеров, обработка ошибок и проверка интеграций. Собственные платёжный процессинг, видеотранскодер и изолирующая среда исполнения не разрабатываются. Это решение для оцениваемого аналога, а не описание внутренней архитектуры Stepik.')
p('За границами: мобильные приложения, офлайн-режим, ML-рекомендации, программы из нескольких курсов, корпоративные классы, LTI, SQL/HTML-задачи, peer review, рассрочки, подарки, автоматизация возвратов и выплат авторам, производство учебного контента и длительное сопровождение. Обращения по возвратам в этой версии обрабатывает оператор вне приложения. Все четыре метода используют именно эти функциональные границы.')
p('Результат разработки — принятая и развёрнутая версия. В WBS включены анализ, дизайн, интеграции, тестирование, документация и выпуск. На одну работу назначается один исполнитель; тесты модуля входят в его реализацию, отдельные работы проверки отвечают за межмодульные сценарии. Неделя — 40 рабочих часов, день — 8. Номинальный человеко-месяц COCOMO — 152 часа [M2].')
p('Авторы сообщили об отсутствии собственного завершённого проекта с известными трудозатратами. Поэтому экспериментальная калибровка PF в пункте 5 задания не выполнена. Раздел UCP содержит полный подсчёт размера и анализ чувствительности к условному PF; это ограничение сохранено в выводах.')
h('2. Функциональные требования',page=True)
p('В колонке «Основание» S — источник из раздела 9; П — проектное допущение. Критерии приёмки относятся к разрабатываемому аналогу. Они делают оценку проверяемой, но не утверждают, что реальный Stepik реализует все технические детали именно так.')
for block in dict.fromkeys(r[1] for r in REQS):
 h(block,2)
 for i,b,n,criterion,source in REQS:
  if b==block:
   p(f'{i}. {n}. {criterion} Основание: [{source}].')
h('Ограничения качества и риски',2)
table(['Область','Принятое допущение / критерий'],[
 ['Нагрузка','До 1 000 одновременно активных сессий и 20 отправок заданий/с; p95 обычного API ≤ 1 с на согласованном стенде. Время внешней проверки учитывается отдельно.'],
 ['Безопасность','Проверка ролей на сервере, хеширование паролей, подписи callback, отсутствие доступа к чужим попыткам. Grader запускает код в готовой изолированной среде.'],
 ['Надёжность','Повторное подтверждение оплаты и результата задания не создаёт второй доступ или дополнительные баллы; резервное копирование и проверка восстановления.'],
 ['Интерфейс','Адаптивная веб-версия, основные сценарии доступны с клавиатуры; состояния ожидания, пустых данных и отказа сервиса.'],
 ['Основной риск','Недооценка интеграции grader и платежей. У работ A10 и A23 расширена пессимистическая оценка; внешние тарифы и эксплуатационные расходы в часах разработки не учитываются.']])
h('3. Наивная оценка',page=True)
p('Для каждой работы заданы оптимистическая O, наиболее вероятная M и пессимистическая P оценки в человеко-часах. Наивный метод суммирует M и предполагает выполнение всех перечисленных работ без дополнительного статистического резерва. Это экспертные оценки, а не измерения.')
table(['ID','Работа','O','M','P'],[[t['id'],t['name'],t['o'],t['m'],t['p']] for t in TASKS])
p(f'Наивная трудоёмкость: H = ΣM = {fmt(NAIVE)} чел·ч. Для одного исполнителя это {fmt(NAIVE/8)} рабочих дней или {fmt(NAIVE/40)} рабочих недель. Деление на размер команды само по себе не учитывает зависимости.')
p('Подробный результат каждой работы и её связь с FR приведены на листе «Работы» книги Excel. A01–A04 создают основу; A05–A26 реализуют функции; A27–A32 проверяют качество и выпускают продукт. Модульные тесты и обычные исправления включены в оценки реализации, чтобы не добавлять их второй раз к итогу.')
h('4. PERT, критический путь и команда',page=True)
p('Для работы i: TEᵢ = (Oᵢ + 4Mᵢ + Pᵢ)/6; σᵢ = (Pᵢ − Oᵢ)/6; Varᵢ = σᵢ². Ожидаемая трудоёмкость H = ΣTEᵢ. При независимости длительностей σH = √ΣVarᵢ. В расчётах сохраняется полная точность, округление применяется только при выводе.')
table(['ID','Предшественники','TE, чел·ч','σ, чел·ч','Var, ч²'],[[t['id'],', '.join(t['deps']) or '—',t['te'],math.sqrt(t['var']),t['var']] for t in TASKS])
p(f'H(PERT) = {fmt(EFF)} чел·ч; σH = {fmt(SIG)} чел·ч. При нормальном приближении 95%-й интервал общей трудоёмкости: H ± 1,96σH = [{fmt(EFF-1.96*SIG)}; {fmt(EFF+1.96*SIG)}] чел·ч. Коррелированные риски могут расширить этот интервал; это не гарантия и не интервал даты релиза.')
h('4.1. Сетевая диаграмма и прямой/обратный проход',2)
p('Используется граф «работа в узле» (AON), связь finish-to-start без задержки. A01 — начало, A32 — завершение. Длительность узла равна TE в рабочих часах при одном исполнителе. ES = max(EF предшественников), EF = ES + TE; LF = min(LS последователей), LS = LF − TE. Полный резерв R = LS − ES. Для начала ES = 0, для завершающей работы LF равно длине критического пути.')
p('Редактируемый полный граф: «Сетевая диаграмма.drawio». Для чтения с увеличением приложены docs/network.svg и docs/network.pdf. Все стрелки построены из того же списка зависимостей, что и таблица CPM.')
picture('network.png','Рисунок 1 — Полная сетевая диаграмма. Красным выделен критический путь.',width=15)
table(['ID','ES','EF','LS','LF','Резерв'],[[t['id'],t['es'],t['ef'],t['ls'],t['lf'],max(0,t['slack'])] for t in TASKS])
p('Критический путь: '+' → '.join(critical)+'.')
p(f'Его длина {fmt(CP)} рабочих часов = {fmt(CP/8)} рабочих дней = {fmt(CP/40)} рабочих недель. Это нижняя граница срока при достаточных ресурсах. Суммарная трудоёмкость остаётся {fmt(EFF)} чел·ч.')
p(f'Для фиксированного критического пути σCP = {fmt(CPSIG)} ч; интервал TE(CP) ± 1,96σCP = [{fmt(CP-1.96*CPSIG)}; {fmt(CP+1.96*CPSIG)}] ч. При изменении длительностей критическим может стать другой путь, поэтому этот интервал не является 95%-м обещанием срока всего проекта.')
h('4.2. Ресурсное расписание и состав команды',2)
p('Построены расписания для 1–8 взаимозаменяемых исполнителей. В момент освобождения ресурса выбирается доступная работа с самым длинным оставшимся путём до конца; при равенстве — меньший ID. Работа не прерывается и не ускоряется назначением нескольких людей. Нижняя граница срока для n человек: max(CP, H/n).')
table(['Человек','Нижняя граница, ч','Полученный срок, ч','Недели','Загрузка, %'],[[n,max(CP,EFF/n),d,d/40,100*EFF/(n*d)] for n,(d,_) in SCHEDULES.items()])
p(f'Критерий выбора: наименьшая из рассмотренных команд, чьё построенное расписание не более чем на 10% длиннее CPM. Рекомендуются {TEAM} разработчика с совмещением ролей; срок {fmt(DURATION/40)} рабочих недель, загрузка {fmt(100*EFF/(TEAM*DURATION))}%. Это рациональный вариант по выбранному критерию, а не доказанный глобальный оптимум задачи ресурсного планирования.')
p(f'Предлагаемое распределение специализаций: два разработчика frontend/backend и {max(TEAM-2,0)} участника с компетенциями full-stack, автоматизации тестирования и DevOps. Дизайн, анализ и документация распределяются по этапам. Предположение о взаимозаменяемости упрощает реальную ситуацию; дефицит нужной компетенции увеличит срок.')
p(f'Для рабочего плана принимается 80% доступности участников на проект. При равномерном снижении доступности срок составляет {fmt(DURATION/0.8/40)} рабочих недель. Дополнительный управленческий резерв 15% даёт {fmt(DURATION/0.8*1.15/40)} недель. Эти проценты — допущения планирования, а не статистический результат PERT; снижение доступности увеличивает срок делением на 0,8.')
picture('schedule.png','Рисунок 2 — Выполнимое расписание по исполнителям.')
table(['ID','Исполнитель','Начало, ч','Окончание, ч'],[[i,w,s,e] for i,w,s,e in sorted(PLAN)])
h('5. Оценка методом функциональных точек',page=True)
p('Выполняется предварительный подсчёт размера нового приложения. Граница охватывает пользовательские функции FR01–FR24; внешние API остаются за ней. ILF — поддерживаемая приложением логическая группа данных, EIF — читаемая группа, поддерживаемая другим приложением. EI — ввод с изменением данных/управления, EO — вывод с вычислением или производным результатом, EQ — запрос без такого вычисления. [M2, раздел 2]')
p('API оплаты, grader, email и хранилище сами по себе не являются EIF. В принятых границах приложение не читает отдельные внешние справочники; ответ сервиса обрабатывается транзакцией, а нужное состояние сохраняется в ILF. Поэтому EIF = 0. Прогресс вычисляется из попыток, а не считается отдельным ILF; наличие отдельной физической таблицы не добавляет функциональные точки.')
p('RET — число логических подгрупп записей, DET — пользовательски различимых полей, FTR — число логических групп, читаемых или изменяемых процессом. Значения предполагаются для проектируемого аналога. Состав групп указан в Excel; это предварительная FPA-оценка, требующая уточнения по спецификации полей и потоков.')
h('5.1. Правила сложности',2)
table(['Группа','Строки матрицы','Столбцы матрицы','Низкая / средняя / высокая'],[
 ['ILF / EIF','RET: 1; 2–5; ≥6','DET: 1–19; 20–50; ≥51','Н Н С / Н С В / С В В'],
 ['EI','FTR: 0–1; 2; ≥3','DET: 1–4; 5–15; ≥16','Н Н С / Н С В / С В В'],
 ['EO / EQ','FTR: 0–1; 2–3; ≥4','DET: 1–5; 6–19; ≥20','Н Н С / Н С В / С В В']])
table(['Тип','Низкая','Средняя','Высокая'],[[k,*v] for k,v in weights.items()])
h('5.2. Логические данные',2)
table(['ID','Группа ILF','RET','DET','Сложность','FP'],[[r[0],r[1],r[3],r[4],r[5],r[6]] for r in FPROWS])
p(f'Сумма ILF = {ILF} FP. EIF = 0 FP. Например, D02 имеет 3 RET и 32 DET: средняя сложность, 10 FP; D01 с 2 RET и 18 DET имеет низкую сложность, 7 FP.')
h('5.3. Транзакции',2)
p('Создание и изменение сущности считаются отдельно там, где различаются элементарный процесс и проверки. Поиск вместе с фильтрами образует один процесс. Табличный отчёт и его CSV-представление объединены при одинаковой обработке. Вариант с разной обработкой потребовал бы отдельного подсчёта. Полный список FTR-групп и связь с FR находятся на листе «FP транзакции».')
table(['ID','Элементарный процесс','Тип','FTR','DET','FP'],[[r[0],r[1],r[2],r[3],r[4],r[6]] for r in TXROWS])
p(f'Транзакции = {TRANS} FP. Невыровненный размер UFP = {ILF} + 0 + {TRANS} = {UFP} FP.')
h('5.4. Выравнивание',2)
p('Для сопоставимости с форматом лабораторной дополнительно рассчитывается традиционный VAF по 14 характеристикам (DI от 0 до 5). Это отдельная учебная поправка; в COCOMO II используется UFP, чтобы не учитывать общие факторы повторно.')
table(['Характеристика','DI'],GSC)
p(f'ΣDI = {sum(r[1] for r in GSC)}; VAF = 0,65 + 0,01 × ΣDI = {fmt(VAF)}; AFP = UFP × VAF = {fmt(AFP)} FP.')
h('6. Трудоёмкость по COCOMO II',page=True)
p('Выбрана COCOMO II.2000 Early Design: архитектура только предполагается и локальной статистики нет. Используется семь агрегированных множителей и пять факторов масштаба. Коэффициенты и формулы взяты из руководства USC [M2, таблицы 62–63]. Уровни факторов — экспертные допущения для оцениваемой команды.')
h('6.1. Размер и факторы',2)
p(f'Предполагаемый прикладной стек — JavaScript на сервере и в браузере; это не утверждение о стеке Stepik. QSM [M1] даёт среднее 47 логических SLOC/FP. Код готовых библиотек и внешних сервисов не включается. SLOC = {UFP} × 47 = {UFP*47}; KSLOC = {fmt(KSLOC)}. Используется UFP = {UFP}; AFP не подставляется.')
table(['Фактор','Значение','Уровень','Обоснование'],SF)
p(f'ΣSF = {fmt(SF_SUM)}; E = 0,91 + 0,01 × ΣSF = {E:.4f}.')
table(['Множитель','Значение','Уровень','Обоснование'],EM)
p(f'EAF = произведение EM = {EAF:.4f}. PM = 2,94 × KSLOC^E × EAF = {fmt(PM)} чел·мес. При 152 чел·ч/месяц трудоёмкость составляет {fmt(COCOMO_H)} чел·ч.')
h('6.2. Срок и чувствительность',2)
p(f'При номинальном SCED: F = 0,28 + 0,2(E − 0,91) = {F:.4f}; TDEV = 3,67 × PM^F = {fmt(TDEV)} календарных месяцев. Средняя численность PM/TDEV = {fmt(PM/TDEV)} человека. Это модельная средняя занятость, а не постоянный штат с первого дня.')
table(['SLOC/FP','KSLOC','PM','Чел·ч','TDEV, мес.'],[[k,UFP*k/1000,2.94*(UFP*k/1000)**E*EAF,2.94*(UFP*k/1000)**E*EAF*152,3.67*(2.94*(UFP*k/1000)**E*EAF)**F] for k in [31,47,63]])
p('Значения 31 и 63 — границы наблюдавшегося диапазона JavaScript в QSM, а не доверительный интервал нашего проекта. Разброс показывает чувствительность к переводу функциональных точек в строки кода. Размер самого FP-подсчёта также неопределён.')
p('COCOMO охватывает инженерные работы и тестирование в определённых моделью фазах; начальная проработка концепции и длительная эксплуатация требуют отдельного согласования границ. В этой работе модель служит независимой проверкой порядка величины, без произвольной добавки «за тестирование». WBS включает ранний анализ и выпуск, поэтому строгое равенство охвата фаз и точности методов не предполагается.')
h('7. Метод Use Case Points',page=True)
p('Акторы классифицируются по интерфейсу, а прецеденты — по числу транзакций взаимодействия: 1–3 → 5, 4–7 → 10, 8 и более → 15. Транзакция — содержательный обмен или завершённая реакция системы, а не отдельный клик. Группировка сценариев является учебным предположением; изменение детализации влияет на UUCW. [M3]')
h('7.1. Акторы и прецеденты',2)
table(['Актор','Вес','Обоснование'],ACTORS)
p(f'UAW = {UAW}. Роли человека имеют вес 3, внешние программные API — 1. Одна учётная запись может выступать в разных ролях; вес относится к типам взаимодействия, а не количеству пользователей.')
table(['ID','Прецедент','FR','Транзакций','Вес'],[[u['id'],u['name'],u['req'],len(u['steps']),u['weight']] for u in UCS])
p(f'UUCW = {UUCW}; UUCP = UUCW + UAW = {UUCW+UAW}. Полные акторы каждого прецедента приведены в книге Excel; ниже перечислены транзакции, по которым получена сложность.')
for u in UCS:
 p(f'{u["id"]}. {u["name"]}: '+ '; '.join(f'{j+1}) {s.lower()}' for j,s in enumerate(u['steps']))+'.')
h('7.2. Технические факторы и окружение',2)
p('Каждому фактору назначено влияние от 0 до 5. TCF = 0,6 + 0,01Σ(wᵢtᵢ); ECF = 1,4 − 0,03Σ(wᵢeᵢ). Отрицательные веса E7 и E8 сохраняются. Оценки окружения относятся к предполагаемой проектной команде, а не выводятся из размера студенческой группы. [M3]')
table(['ID','Технический фактор','Вес','Оценка','Вклад'],[[i,n,w,v,w*v] for i,n,w,v,x in TCF])
table(['ID','Фактор окружения','Вес','Оценка','Вклад'],[[i,n,w,v,w*v] for i,n,w,v,x in ECF])
p(f'TF = {fmt(TF)}; TCF = {TC:.4f}; EF = {fmt(EF)}; ECF = {EC:.4f}. UCP = ({UUCW} + {UAW}) × {TC:.4f} × {EC:.4f} = {fmt(UCP)}. Единица UCP — размер, а не человеко-часы.')
h('7.3. Фактор продуктивности: доступные данные',2)
p('Собственного завершённого проекта с известными трудозатратами у авторов нет. Поэтому вычислить эмпирический PF = Href / UCPref по условию задания невозможно. Часы, функции и результаты вымышленного проекта не подставляются. Эта часть задания остаётся ограничением работы.')
p('Для завершения калибровки в будущем нужны: название реального проекта; перечень акторов и транзакций; TCF и ECF; фактические суммарные человеко-часы всех участников в согласованных границах. Затем UCPref = (UUCWref + UAWref) × TCFref × ECFref; PF = Href/UCPref; HStepik = UCP × PF. При неполном учёте тестирования или дизайна PF нельзя без поправок переносить на полный жизненный цикл.')
p('Показан сценарный расчёт H = UCP × PF. Значение 20 чел·ч/UCP используется как исторический ориентир метода, обсуждаемый в [M4]; 10 и 30 выбраны для анализа чувствительности. Ни одно из них не является измеренным PF авторов. Добавочная константа не вводится, чтобы произвольно не учитывать одни и те же работы дважды.')
table(['Условный PF, чел·ч/UCP','Трудоёмкость, чел·ч','Статус'],[[p,UCP*p,'Сценарий без локальной калибровки'] for p in [10,20,30]])
h('8. Сравнение результатов и вывод',page=True)
table(['Метод','Трудоёмкость, чел·ч','Что означает результат'],[
 ['Наивный',NAIVE,'Сумма наиболее вероятных оценок WBS.'],['PERT',EFF,'Ожидаемая сумма всех работ; без подмены критическим путём.'],['FP + COCOMO II',COCOMO_H,'Параметрическая оценка при выбранных SF/EM и 47 SLOC/FP.'],['UCP, условный PF=20',UCP*20,'Сценарный ориентир; не результат калибровки на собственном проекте.']])
p(f'PERT превышает наивную оценку на {fmt((EFF/NAIVE-1)*100)}%: пессимистические исходы дают асимметрию для сложных интеграций. Разница не доказывает точность одного метода: оба используют один экспертный WBS и одинаковые границы.')
p(f'Ресурсный граф даёт минимальную длительность {fmt(CP/40)} рабочих недель, а выбранное расписание для {TEAM} исполнителей — {fmt(DURATION/40)} недели при полной доступности. С учётом доступности 80% и отдельного резерва 15% плановый ориентир — {math.ceil(DURATION/0.8*1.15/40)} рабочих недель. Эти сроки нельзя записывать в колонку человеко-часов.')
p(f'COCOMO даёт {fmt(COCOMO_H/EFF)} от трудоёмкости PERT. Расхождение обусловлено отраслевыми коэффициентами, предварительным FP-подсчётом, предположением о новом прикладном коде и различием границ фаз. Оно указывает на необходимость проверки WBS и интеграционных прототипов; отраслевое происхождение коэффициентов само по себе не делает оценку достовернее для конкретной команды.')
p(f'UCP даёт размер {fmt(UCP)} и линейную зависимость трудоёмкости от PF. Без исторических данных нельзя выбрать единственную подтверждённую трудоёмкость. Условные значения помогают увидеть чувствительность, но не заменяют требуемый опыт завершённого проекта.')
p('Вывод: сформированы требования к аналогу Stepik, выполнены WBS, PERT и CPM, построен ресурсный план, подсчитаны функциональные точки, применены COCOMO II и UCP. Для первого планирования полезен прозрачный WBS/PERT с резервом, а COCOMO и UCP служат независимыми проверками. После прототипа оплаты и проверки кода следует уточнить O/M/P и собрать фактическую скорость команды. Экспериментальная калибровка PF по собственному проекту не выполнена из-за отсутствия исходных данных; её необходимо дополнить для полного выполнения соответствующего пункта задания.')
h('9. Источники',page=True)
p('Дата обращения ко всем веб-источникам: '+DATE+'. Требования пересказаны, а не скопированы. Ссылки S подтверждают функции Stepik, M — используемые методы; численные оценки проекта являются самостоятельными допущениями.')
for i,n,url,x in SOURCES: p(f'[{i}] {n}. {url} — {x}')
p('Образец структуры: предоставленные пользователем файлы ~/epi1_max (README, отчёт, Excel и draw.io).')
doc.core_properties.title='Лабораторная работа № 1 — Stepik'
doc.core_properties.author='; '.join(AUTHORS)
doc.core_properties.subject='Экономика программной инженерии'
doc.save(DOCS/'report.docx')
(DOCS/'report.md').write_text('\n'.join(md),encoding='utf-8')
results=dict(naive_hours=NAIVE,pert_hours=EFF,pert_sigma=SIG,critical_path=critical,critical_hours=CP,team=TEAM,schedule_hours=DURATION,planned_weeks=math.ceil(DURATION/.8*1.15/40),ilf=ILF,transactions_fp=TRANS,ufp=UFP,afp=AFP,ksloc=KSLOC,cocomo_pm=PM,cocomo_hours=COCOMO_H,cocomo_months=TDEV,uucw=UUCW,uaw=UAW,tcf=TC,ecf=EC,ucp=UCP,ucp_hours_pf20=UCP*20,pf_calibrated=False)
(DOCS/'calculations.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(results,ensure_ascii=False,indent=2))
